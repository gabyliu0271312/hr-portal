from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException

from app.reports import report_service


@pytest.mark.parametrize("field", ["execution_version", "execution_binding"])
def test_ordinary_report_input_rejects_execution_controls(field):
    from app.reports.router import ReportIn, RunOverrides
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        ReportIn.model_validate({"name": "report", "dataset_id": 1, field: "legacy_v1"})
    with pytest.raises(ValidationError):
        RunOverrides.model_validate({field: "legacy_v1"})


def _report():
    report = SimpleNamespace(id=1, dataset_id=2, owner_id=7, scope_strategy=None,
        config={"columns": ["r.amount"], "aggregate": True, "aggregations": {"r.amount": "sum"}},
        execution_version="authorized_v1")
    report.execution_binding = {"fingerprint": report_service.execution_fingerprint(report),
        "approved_by": 7, "approved_at": "2026-10-21T00:00:00+00:00", "reason": "review",
        "contract_id": 11, "contract_version": 1}
    return report


@pytest.mark.asyncio
async def test_authorized_execution_requires_binding_and_never_falls_back(monkeypatch):
    report = _report()
    report.execution_binding = {}
    legacy = AsyncMock()
    monkeypatch.setattr("app.reports.sql_builder.run_dataset_query", legacy)
    audit = AsyncMock()
    monkeypatch.setattr("app.reports.audit.record_authorized_scope_audit", audit)
    with pytest.raises(HTTPException) as error:
        await report_service.run_audited_dataset_query(
            channel="run", report=report, report_id=1, dataset_id=2, user=SimpleNamespace(id=7), db=object())
    assert error.value.status_code == 409
    assert error.value.detail["code"] == "execution_binding_invalid"
    legacy.assert_not_awaited()
    assert audit.await_args.kwargs["query_status"] == "failed"


@pytest.mark.asyncio
async def test_legacy_execution_version_continues_using_legacy_query(monkeypatch):
    report = SimpleNamespace(id=1, execution_version="legacy_v1")
    legacy = AsyncMock(return_value=([], [], 0))
    audit = AsyncMock()
    monkeypatch.setattr("app.reports.sql_builder.run_dataset_query", legacy)
    monkeypatch.setattr("app.reports.audit.record_scope_audit", audit)
    result = await report_service.run_audited_dataset_query(
        channel="run", report=report, report_id=1, dataset_id=2, user=SimpleNamespace(id=7), db=object())
    assert result == ([], [], 0)
    legacy.assert_awaited_once()
    assert "report" not in legacy.await_args.kwargs
    audit.assert_awaited_once()


@pytest.mark.asyncio
@pytest.mark.parametrize("channel", ["push", "push_columns", "scheduled", "dimension_preview", "model_preview", "export_rows"])
async def test_authorized_unsupported_channels_reject_before_query(channel, monkeypatch):
    gate = AsyncMock()
    monkeypatch.setattr(report_service, "require_authorized_report_readiness", gate)
    with pytest.raises(HTTPException) as error:
        await report_service.run_authorized_report_query(report=_report(), channel=channel, dataset_id=2)
    assert error.value.detail["code"] == "authorized_channel_not_supported"
    gate.assert_not_awaited()


@pytest.mark.asyncio
@pytest.mark.parametrize("change", ["config", "dataset_id", "owner_id", "id", "scope_strategy"])
async def test_authorized_binding_invalidates_on_report_identity_change(change):
    report = _report()
    setattr(report, change, {} if change == "config" else 999 if change != "scope_strategy" else "cc_first")
    with pytest.raises(HTTPException) as error:
        await report_service.run_authorized_report_query(report=report, channel="run", dataset_id=report.dataset_id)
    assert error.value.detail["code"] == "execution_binding_invalid"


@pytest.mark.asyncio
async def test_current_readiness_cannot_be_replaced_by_published_or_active_contract(monkeypatch):
    from app.datasets.models import DataSet
    dataset = DataSet(id=2, created_by=7, status="published", is_active=True)
    db = SimpleNamespace(get=AsyncMock(return_value=dataset), scalar=AsyncMock(return_value=None))
    release = SimpleNamespace(eligible=True, fields_eligible=True, authorization_verified=True, conversion_verified=False)
    readiness = SimpleNamespace(eligible=True, runtime_enabled=True)
    monkeypatch.setattr("app.warehouse.service.release_gate.check_dataset_release", AsyncMock(return_value=release))
    monkeypatch.setattr("app.datasets.authorization_contract.execution_readiness", AsyncMock(return_value=readiness))
    with pytest.raises(HTTPException) as error:
        await report_service.require_authorized_report_readiness(_report(), SimpleNamespace(id=7), db)
    assert error.value.detail["code"] == "authorized_execution_not_ready"


@pytest.mark.asyncio
@pytest.mark.parametrize("channel", ["run", "csv", "xlsx"])
async def test_explicit_authorized_dispatch_never_calls_legacy(channel, monkeypatch):
    gate = AsyncMock(return_value=SimpleNamespace(contract_id=11, contract_version=1))
    execute = AsyncMock(return_value=([{"code": "r.amount"}], [{"r.amount": "600"}], 1))
    legacy = AsyncMock()
    monkeypatch.setattr(report_service, "require_authorized_report_readiness", gate)
    monkeypatch.setattr(report_service, "_execute_authorized_summary", execute)
    monkeypatch.setattr("app.reports.sql_builder.run_dataset_query", legacy)
    monkeypatch.setattr("app.reports.audit.record_authorized_scope_audit", AsyncMock())
    result = await report_service.run_audited_dataset_query(channel=channel, report=_report(),
        report_id=1, dataset_id=2, user=SimpleNamespace(id=7), db=object())
    assert result[2] == 1
    execute.assert_awaited_once()
    legacy.assert_not_awaited()


@pytest.mark.asyncio
async def test_authorized_binding_cannot_be_skipped_by_omitting_report(monkeypatch):
    legacy = AsyncMock()
    monkeypatch.setattr("app.reports.sql_builder.run_dataset_query", legacy)
    monkeypatch.setattr("app.reports.audit.record_authorized_scope_audit", AsyncMock())
    db = SimpleNamespace(get=AsyncMock(return_value=_report()))
    with pytest.raises(HTTPException) as error:
        await report_service.run_audited_dataset_query(channel="push", report_id=1, dataset_id=2, db=db)
    assert error.value.detail["code"] == "authorized_channel_not_supported"
    legacy.assert_not_awaited()


@pytest.mark.asyncio
async def test_authorized_report_realtime_export_is_blocked_before_sql():
    from app.reports.sql_builder import assess_report_realtime_deployability
    assessment = await assess_report_realtime_deployability(_report(), None)
    assert assessment["supported"] is False


@pytest.mark.postgres_acceptance
@pytest.mark.asyncio
async def test_postgres_binding_api_keeps_legacy_when_full_readiness_is_missing(standard_catalog_api):
    from sqlalchemy import text
    from app.reports.report_service import execution_fingerprint
    from app.reports.models import Report
    from app.users.models import Menu, RoleMenu
    client, engine, factory, app = standard_catalog_api
    from app.reports.router import router
    app.include_router(router, prefix="/api/v1")
    async with engine.begin() as conn:
        await conn.execute(Menu.__table__.insert().values(id=8, code="report.list", label="Reports"))
        await conn.execute(RoleMenu.__table__.insert().values(id=8, role_id=1, menu_id=8, can_view=True, can_update=True))
        await conn.execute(Report.__table__.insert().values(id=1, name="isolated report", table_name="", dataset_id=4,
            config={"columns": ["current.employee_no"]}, owner_id=7))
    response = await client.get("/api/v1/reports/1/execution-binding")
    assert response.status_code == 200, response.text
    assert response.json()["execution_version"] == "legacy_v1"
    payload = {"expected_fingerprint": response.json()["fingerprint"], "contract_id": 1, "contract_version": 1, "reason": "review"}
    response = await client.post("/api/v1/reports/1/execution-binding", json=payload)
    assert response.status_code == 409, response.text
    assert response.json()["detail"]["code"] == "authorized_execution_not_ready"
    async with engine.connect() as conn:
        assert (await conn.execute(text("SELECT execution_version,execution_binding FROM reports WHERE id=1"))).one() == ("legacy_v1", None)
    response = await client.post("/api/v1/reports/1/execution-binding", json={**payload, "approved_by": 7})
    assert response.status_code == 422


@pytest.mark.postgres_acceptance
@pytest.mark.asyncio
async def test_postgres_execution_migration_preserves_legacy_and_blocks_unsafe_downgrade(standard_catalog_api):
    from importlib.util import spec_from_file_location, module_from_spec
    from pathlib import Path
    from sqlalchemy import text
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    _, engine, _, _ = standard_catalog_api
    spec = spec_from_file_location("report_binding_migration", Path(__file__).parents[1] / "alembic/versions/0267_report_execution_binding.py")
    migration = module_from_spec(spec)
    spec.loader.exec_module(migration)
    def apply(conn, action):
        migration.op = Operations(MigrationContext.configure(conn))
        getattr(migration, action)()
    async with engine.begin() as conn:
        await conn.execute(text("INSERT INTO reports(id,name,table_name,dataset_id,config,owner_id,is_published,visibility,run_count) VALUES(1,'preserved','',4,'{}',7,false,'private',0)"))
        await conn.run_sync(lambda sync: apply(sync, "downgrade"))
        await conn.run_sync(lambda sync: apply(sync, "upgrade"))
        assert (await conn.execute(text("SELECT name,execution_version,execution_binding FROM reports WHERE id=1"))).one() == ("preserved", "legacy_v1", None)
        await conn.execute(text("UPDATE reports SET execution_version='authorized_v1',execution_binding='{}' WHERE id=1"))
        with pytest.raises(RuntimeError, match="safely retired"):
            await conn.run_sync(lambda sync: apply(sync, "downgrade"))
        assert await conn.scalar(text("SELECT execution_version FROM reports WHERE id=1")) == "authorized_v1"


@pytest.mark.postgres_acceptance
@pytest.mark.asyncio
@pytest.mark.parametrize("can_view,can_update,warehouse_update,expected", [
    (True, True, True, 200), (True, True, False, 403), (True, False, True, 403),
])
async def test_postgres_binding_route_requires_both_operations(standard_catalog_api, can_view, can_update, warehouse_update, expected):
    from app.users.models import Menu, RoleMenu
    from app.reports.models import Report
    from app.reports.router import router
    client, engine, _, app = standard_catalog_api
    app.include_router(router, prefix="/api/v1")
    async with engine.begin() as conn:
        await conn.execute(Menu.__table__.insert().values(id=8, code="report.list", label="Reports"))
        await conn.execute(RoleMenu.__table__.insert().values(id=8, role_id=1, menu_id=8, can_view=can_view, can_update=can_update))
        await conn.execute(RoleMenu.__table__.update().where(RoleMenu.id == 1).values(can_update=warehouse_update))
        await conn.execute(Report.__table__.insert().values(id=1, name="permission report", table_name="", dataset_id=4, config={}, owner_id=7))
    response = await client.get("/api/v1/reports/1/execution-binding")
    assert response.status_code == expected
    if expected == 200:
        response = await client.post("/api/v1/reports/1/execution-binding", json={
            "expected_fingerprint": "0" * 64, "contract_id": 1, "contract_version": 1, "reason": "review"})
        assert response.status_code == 409
        assert response.json()["detail"]["code"] == "execution_binding_conflict"


@pytest.mark.postgres_acceptance
@pytest.mark.asyncio
async def test_postgres_binding_transition_is_atomic_under_synthetic_readiness(standard_catalog_api, monkeypatch):
    from sqlalchemy import text
    from app.users.models import Menu, RoleMenu
    from app.reports.models import Report
    from app.reports.router import router
    client, engine, _, app = standard_catalog_api
    app.include_router(router, prefix="/api/v1")
    monkeypatch.setattr(report_service, "require_authorized_report_readiness",
        AsyncMock(return_value=SimpleNamespace(contract_id=11, contract_version=1)))
    async with engine.begin() as conn:
        await conn.execute(Menu.__table__.insert().values(id=8, code="report.list", label="Reports"))
        await conn.execute(RoleMenu.__table__.insert().values(id=8, role_id=1, menu_id=8, can_view=True, can_update=True))
        await conn.execute(Report.__table__.insert().values(id=1, name="binding report", table_name="", dataset_id=4, config={}, owner_id=7))
    response = await client.get("/api/v1/reports/1/execution-binding")
    payload = {"expected_fingerprint": response.json()["fingerprint"], "contract_id": 11, "contract_version": 1, "reason": "private reason"}
    async with engine.begin() as conn:
        await conn.execute(text("ALTER TABLE system_logs ADD CONSTRAINT forced_failure CHECK (action <> 'report_execution_binding')"))
    with pytest.raises(Exception):
        await client.post("/api/v1/reports/1/execution-binding", json=payload)
    async with engine.begin() as conn:
        assert await conn.scalar(text("SELECT execution_version FROM reports WHERE id=1")) == "legacy_v1"
        await conn.execute(text("ALTER TABLE system_logs DROP CONSTRAINT forced_failure"))
    response = await client.post("/api/v1/reports/1/execution-binding", json=payload)
    assert response.status_code == 200, response.text
    async with engine.connect() as conn:
        assert await conn.scalar(text("SELECT execution_version FROM reports WHERE id=1")) == "authorized_v1"
        logs = (await conn.execute(text("SELECT metadata_json FROM system_logs WHERE action='report_execution_binding'"))).scalars().all()
        assert len(logs) == 1 and "private reason" not in str(logs)


@pytest.mark.postgres_acceptance
@pytest.mark.asyncio
async def test_postgres_shared_dataset_external_target_blocks_authorized_readiness(standard_catalog_api):
    from app.reports.models import Report
    from app.push.models import PushTarget
    client, engine, factory, _ = standard_catalog_api
    async with engine.begin() as conn:
        await conn.execute(Report.__table__.insert().values(
            id=5, name="shared report", table_name="", dataset_id=4, config={}, owner_id=7))
        await conn.execute(PushTarget.__table__.insert().values(
            id=1, source_table="report:5", source_type="table", name="legacy target", push_type="db_snapshot"))
    report = SimpleNamespace(id=6, dataset_id=4)
    async with factory() as db:
        with pytest.raises(HTTPException) as error:
            await report_service.require_authorized_report_readiness(report, SimpleNamespace(id=7), db)
        assert error.value.detail["code"] == "external_exports_not_isolated"


@pytest.mark.postgres_acceptance
@pytest.mark.asyncio
async def test_postgres_run_csv_xlsx_use_real_summary_under_synthetic_admission(postgres_d_fact, monkeypatch):
    import csv
    import io
    import httpx
    from datetime import datetime, timezone
    from fastapi import FastAPI
    from openpyxl import load_workbook
    from app.core.db import get_session
    from app.core.deps import current_user
    from app.datasets.models import DataProductAuthorizationContract, DataSetAcl, DatasetCalculatedField
    from app.reports.models import Report
    from app.reports.router import router
    from app.users.models import Menu, RoleMenu
    from app.warehouse.models import WarehouseQualityRule, WarehouseQualityStatus
    from tests.test_d_execution import _s_contract

    engine, factory, _ = postgres_d_fact
    contract = _s_contract()
    config = {"columns": ["allocation.cc_code", "allocation.amount"], "aggregate": True,
              "aggregations": {"allocation.amount": "sum"}}
    async with engine.begin() as conn:
        for model in (DatasetCalculatedField, WarehouseQualityRule, WarehouseQualityStatus):
            await conn.run_sync(lambda sync, model=model: model.__table__.create(sync, checkfirst=True))
        await conn.execute(Menu.__table__.insert().values(id=8, code="report.list", label="Reports"))
        await conn.execute(RoleMenu.__table__.insert().values(id=8, role_id=1, menu_id=8,
            can_view=True, can_export=True))
        await conn.execute(DataSetAcl.__table__.insert().values(id=7, dataset_id=7, user_id=7))
        await conn.execute(DataProductAuthorizationContract.__table__.insert().values(
            id=12, dataset_id=7, contract_version=1, status="active", created_by=7,
            contract_json=contract.definition.model_dump(), model_snapshot=contract.model_snapshot,
            approved_by=7, approved_at=datetime.now(timezone.utc), approval_reason="synthetic contract"))
        report = Report(id=1, name="summary", table_name="", dataset_id=7, owner_id=7,
                        config=config, scope_strategy=None, execution_version="authorized_v1")
        binding = {"fingerprint": report_service.execution_fingerprint(report), "contract_id": 12,
                   "contract_version": 1, "approved_by": 7,
                   "approved_at": datetime.now(timezone.utc).isoformat(), "reason": "synthetic admission"}
        await conn.execute(Report.__table__.insert().values(id=1, name="summary", table_name="",
            dataset_id=7, owner_id=7, config=config, execution_version="authorized_v1", execution_binding=binding))
    gate = AsyncMock(return_value=SimpleNamespace(contract_id=12, contract_version=1))
    monkeypatch.setattr(report_service, "require_authorized_report_readiness", gate)
    legacy = AsyncMock(side_effect=AssertionError("authorized result must not call legacy"))
    monkeypatch.setattr("app.reports.sql_builder.run_dataset_query", legacy)
    monkeypatch.setattr("app.reports.router.record_report_audit", AsyncMock())
    monkeypatch.setattr("app.reports.audit.record_authorized_scope_audit", AsyncMock())
    monkeypatch.setattr("app.automation.events.publish_event", AsyncMock())
    monkeypatch.setattr("app.core.db.get_session_factory", lambda: factory)
    monkeypatch.setattr("app.reports.router._prepare_xlsx_export_rows",
        lambda *args: (_ for _ in ()).throw(AssertionError("no second rounding pass")))
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    async def session():
        async with factory() as db:
            yield db
    app.dependency_overrides[get_session] = session
    app.dependency_overrides[current_user] = lambda: SimpleNamespace(id=7)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post("/api/v1/reports/1/run?page=1&page_size=1")
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["total"] == 2 and len(body["items"]) == 1
        assert body["warnings"]
        response = await client.get("/api/v1/reports/1/export.csv")
        assert response.status_code == 200, response.text
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        assert len(rows) == 3 and sum(float(row[1].replace(",", "")) for row in rows[1:]) == 600
        response = await client.get("/api/v1/reports/1/export.xlsx")
        assert response.status_code == 200, response.text
        workbook = load_workbook(io.BytesIO(response.content))
        matrix = list(workbook.active.values)
        assert len(matrix) == 3 and sum(row[1] for row in matrix[1:]) == 600
        response = await client.post("/api/v1/reports/1/run", json={"filters": [{"column": "allocation.cc_code", "value": "CC1"}]})
        assert response.status_code == 422
        assert response.json()["detail"]["code"] == "authorized_runtime_filters_unsupported"
        response = await client.get("/api/v1/reports/1/export.csv", params={"runtime_filters": '[{"column":"allocation.cc_code","value":"CC1"}]'})
        assert response.status_code == 422
        response = await client.get("/api/v1/reports/1/export.xlsx", params={"runtime_filters": '[{"column":"allocation.cc_code","value":"CC1"}]'})
        assert response.status_code == 422
    assert gate.await_count == 3
    legacy.assert_not_awaited()


@pytest.mark.postgres_acceptance
@pytest.mark.asyncio
@pytest.mark.parametrize("authorization,expected", [("org", 400), ("cc", 300), ("both", 600), ("none", None)])
async def test_postgres_real_conversion_contract_binding_run_count_csv_xlsx(postgres_d_fact, monkeypatch, authorization, expected):
    import csv
    import io
    import httpx
    from fastapi import FastAPI
    from openpyxl import load_workbook
    from sqlalchemy import MetaData, Table, select, text
    from app.core.db import get_session
    from app.core.deps import current_user
    from app.data import dynamic_loader, models
    from app.datasets.models import DataSetAcl, DatasetCalculatedField
    from app.reports.models import Report
    from app.reports.router import router
    from app.scopes.models import UserScopeTag
    from app.users.models import Menu, RoleMenu
    from app.warehouse.models import StandardizationRule, WarehouseQualityRule, WarehouseQualityStatus
    from app.warehouse.router import router as warehouse_router
    from app.warehouse.service.standardization import StandardizationRuleService

    engine, factory, _ = postgres_d_fact
    async with engine.begin() as conn:
        for model in (DatasetCalculatedField, WarehouseQualityRule, WarehouseQualityStatus):
            await conn.run_sync(lambda sync, model=model: model.__table__.create(sync, checkfirst=True))
        await conn.execute(text("CREATE TABLE allocation_source (fact_id bigint, employee_no text, amount numeric, cc_code text, department text)"))
        await conn.execute(text("INSERT INTO allocation_source SELECT fact_id,employee_no,amount,cc_code,department FROM allocation_fact"))
        await conn.execute(models.RegisteredTable.__table__.insert().values(
            id=33, table_name="allocation_source", table_label="Source", warehouse_layer="ODS"))
        columns = (await conn.execute(select(models.TableColumn.__table__).where(
            models.TableColumn.table_name == "allocation_fact"))).mappings().all()
        for column in columns:
            await conn.execute(models.TableColumn.__table__.insert().values(
                table_name="allocation_source", column_code=column["column_code"],
                column_label=column["column_label"], data_type=column["data_type"],
                is_visible=True, scope_role=column["scope_role"]))
        await conn.execute(models.TableColumn.__table__.insert().values(
            table_name="allocation_fact", column_code="id", column_label="Internal identity", data_type="number"))
        await conn.execute(StandardizationRule.__table__.insert().values(
            asset_type="table", asset_code="allocation_source", rule_type="rename",
            source_field="amount", target_field="amount", rule_config={}, enabled=True))
        await conn.execute(Menu.__table__.insert().values(id=8, code="report.list", label="Reports"))
        await conn.execute(RoleMenu.__table__.insert().values(id=8, role_id=1, menu_id=8,
            can_view=True, can_update=True, can_export=True))
        await conn.execute(DataSetAcl.__table__.insert().values(id=7, dataset_id=7, user_id=7))
        await conn.execute(Report.__table__.insert().values(id=1, name="real summary", table_name="", dataset_id=7,
            owner_id=7, config={"columns": ["allocation.cc_code", "allocation.amount"], "aggregate": True,
                                "aggregations": {"allocation.amount": "sum"}}))

    async def reflect_in_schema(db, table_name):
        reflected = await db.run_sync(lambda sync: Table(table_name, MetaData(), autoload_with=sync.connection()))
        return dynamic_loader._make_model_from_table(table_name, reflected)

    monkeypatch.setattr(dynamic_loader, "reflect_source_table_model", reflect_in_schema)
    async with factory() as db:
        refreshed = await StandardizationRuleService(db).execute_full(
            asset_code="allocation_source", target_table="allocation_fact")
        assert "error" not in refreshed, refreshed
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    app.include_router(warehouse_router, prefix="/api/v1")

    async def session():
        async with factory() as db:
            yield db

    app.dependency_overrides[get_session] = session
    app.dependency_overrides[current_user] = lambda: SimpleNamespace(id=7)
    monkeypatch.setattr("app.core.db.get_session_factory", lambda: factory)
    monkeypatch.setattr("app.reports.audit.get_session_factory", lambda: factory)
    monkeypatch.setattr("app.reports.router.record_report_audit", AsyncMock())
    monkeypatch.setattr("app.automation.events.publish_event", AsyncMock())
    legacy = AsyncMock(side_effect=AssertionError("authorized execution must not fall back"))
    monkeypatch.setattr("app.reports.sql_builder.run_dataset_query", legacy)
    model_base = "/api/v1/warehouse/models/7"
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
        created = await client.post(model_base + "/authorization-contracts", json={
            "expected_version": 0,
            "facts": [{"alias": "allocation", "table_name": "allocation_fact", "fact_key": "id",
                       "employee_key": "employee_no", "cost_center_key": "cc_code",
                       "allow_cost_center_summary": True, "role": "allocation_fact"}],
            "output_fields": ["cc_code", "department", "amount"], "dimensions": ["cc_code", "department"],
            "measures": ["amount"], "internal_fields": ["allocation.id", "allocation.employee_no"],
        })
        assert created.status_code == 201, created.text
        contract_id = created.json()["id"]
        approved = await client.post(f"{model_base}/authorization-contracts/{contract_id}/approve",
                                     json={"expected_revision": 0, "reason": "real contract review"})
        assert approved.status_code == 200, approved.text
        before = await client.get(model_base + "/authorization-execution-readiness")
        assert before.status_code == 200 and before.json()["runtime_enabled"] is False
        conversion = await client.get(model_base + "/conversion-readiness")
        assert conversion.status_code == 200, conversion.text
        approved = await client.post(model_base + "/conversion-readiness/approve", json={
            "expected_run_id": conversion.json()["run_id"], "expected_contract_id": contract_id,
            "expected_dataset_version": 1, "reason": "real conversion review"})
        assert approved.status_code == 200, approved.text
        ready = await client.get(model_base + "/authorization-execution-readiness")
        assert ready.status_code == 200 and ready.json()["runtime_enabled"] is True, ready.text
        assert ready.json()["paths"][0]["detail_mode"] == "not_allowed"
        from app.datasets.d_execution import SExecutionBindingIn, compile_allocation_fact_s_select_for_dataset
        from app.warehouse.schemas import ExecutionPlanIn
        async with factory() as db:
            statement, issues = await compile_allocation_fact_s_select_for_dataset(
                db=db, user=SimpleNamespace(id=7), request=SExecutionBindingIn(
                    dataset_id=7, contract_id=contract_id, contract_version=1,
                    plan=ExecutionPlanIn(output_fields=["cc_code", "amount"], group_by=["cc_code"],
                                         aggregations={"amount": "sum"})))
            assert statement is not None and issues == []
            assert sum(row.amount for row in (await db.execute(statement)).all()) == 600
        state = await client.get("/api/v1/reports/1/execution-binding")
        bound = await client.post("/api/v1/reports/1/execution-binding", json={
            "expected_fingerprint": state.json()["fingerprint"], "contract_id": contract_id,
            "contract_version": 1, "reason": "real isolated trial"})
        assert bound.status_code == 200, bound.text
        retained = {"org": [1], "cc": [2], "both": [1, 2], "none": []}[authorization]
        async with engine.begin() as conn:
            await conn.execute(UserScopeTag.__table__.delete().where(UserScopeTag.tag_id.not_in(retained)))
        response = await client.post("/api/v1/reports/1/run?page=1&page_size=1")
        if expected is None:
            assert response.status_code in {403, 409}, response.text
            for suffix in ("csv", "xlsx"):
                assert (await client.get(f"/api/v1/reports/1/export.{suffix}")).status_code in {403, 409}
        else:
            assert response.status_code == 200, response.text
            assert response.json()["total"] == (1 if authorization == "cc" else 2)
            assert len(response.json()["items"]) == 1
            exported = await client.get("/api/v1/reports/1/export.csv")
            assert exported.status_code == 200, exported.text
            rows = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))
            assert sum(float(row[1]) for row in rows[1:]) == expected
            exported = await client.get("/api/v1/reports/1/export.xlsx")
            assert exported.status_code == 200, exported.text
            matrix = list(load_workbook(io.BytesIO(exported.content)).active.values)
            assert sum(row[1] for row in matrix[1:]) == expected
            if authorization == "both":
                from app.push.models import PushTarget
                async with engine.begin() as conn:
                    await conn.execute(PushTarget.__table__.insert().values(
                        id=88, source_type="table", source_id="allocation_fact", source_table="allocation_fact",
                        name="isolated external target", push_type="api_expose", is_active=True))
                assert (await client.post("/api/v1/reports/1/run")).status_code == 409
                async with engine.begin() as conn:
                    await conn.execute(text("UPDATE push_targets SET is_active=false WHERE id=88"))
                    await conn.execute(text("DELETE FROM dataset_acl WHERE dataset_id=7 AND user_id=7"))
                assert (await client.post("/api/v1/reports/1/run")).status_code == 403
                async with engine.begin() as conn:
                    await conn.execute(DataSetAcl.__table__.insert().values(id=7, dataset_id=7, user_id=7))
                    await conn.execute(text("UPDATE standardization_rules SET enabled=false WHERE asset_code='allocation_source'"))
                assert (await client.post("/api/v1/reports/1/run")).status_code == 409
                async with engine.begin() as conn:
                    await conn.execute(text("UPDATE standardization_rules SET enabled=true WHERE asset_code='allocation_source'"))
                assert (await client.post("/api/v1/reports/1/run")).status_code == 200
            async with engine.begin() as conn:
                await conn.execute(text("UPDATE table_columns SET is_visible=false WHERE table_name='allocation_fact' AND column_code='amount'"))
            denied = await client.post("/api/v1/reports/1/run")
            assert denied.status_code in {403, 409, 422}, denied.text
            assert "items" not in denied.json()
        legacy.assert_not_awaited()


from tests.test_d_execution import postgres_d_fact
from tests.test_scope_classification import postgres_classification_schema
from tests.test_standard_field_catalog import postgres_standard_catalog, standard_catalog_api
