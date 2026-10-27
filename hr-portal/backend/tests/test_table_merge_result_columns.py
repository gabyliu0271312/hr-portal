import asyncio
import io
import json

import openpyxl
import pytest

from app.table_tools import engine
from app.table_tools import router as table_routes
from app.table_tools.models import MergeDwdRelation, MergeTemplate


def _template() -> MergeTemplate:
    template = MergeTemplate(
        id=1,
        name="template",
        merge_keys=["证件号码"],
        std_fields=["金额"],
        aggregate="sum",
    )
    template.dwd_relations.append(MergeDwdRelation(
        id=1,
        template_id=1,
        name="DWD 关联",
        dataset_id=16,
        left_fields=["证件号码"],
        right_fields=["dwd_current.id_number"],
        select_fields=["dwd_current.corporate_email"],
        enabled=True,
    ))
    return template


def test_run_template_merge_places_dwd_columns_before_source_and_returns_labels(monkeypatch):
    template = _template()

    async def functions(_db):
        return {}

    monkeypatch.setattr(table_routes, "executable_functions", functions)
    monkeypatch.setattr(table_routes.engine, "run_merge", lambda *_args: {
        "rows": [{"证件号码": "110", "金额": 1, "来源": "源表"}],
        "columns": ["证件号码", "金额", "来源"],
        "recognize_log": [], "anomalies": [], "stats": {},
        "key_mapping_stats": {}, "raw_key_traces": [],
    })

    async def apply_relation(rows, *_args):
        return [{**row, "dwd_current.corporate_email": "a@example.com"} for row in rows], []

    async def list_fields(*_args):
        return [{"code": "dwd_current.corporate_email", "label": "企业邮箱"}]

    monkeypatch.setattr(table_routes, "apply_dwd_relation", apply_relation)
    monkeypatch.setattr(table_routes, "list_dwd_fields_by_dataset", list_fields)

    result = asyncio.run(table_routes._run_template_merge(
        template, [], object(), object(), relations=list(template.dwd_relations)
    ))

    assert result["columns"] == ["证件号码", "金额", "dwd_current.corporate_email", "来源"]
    assert result["column_labels"] == {"dwd_current.corporate_email": "企业邮箱"}


def test_rows_to_xlsx_uses_codes_for_values_and_labels_for_headers():
    content = engine.rows_to_xlsx(
        ["dwd_current.corporate_email", "dwd_current.personal_email"],
        [{
            "dwd_current.corporate_email": "work@example.com",
            "dwd_current.personal_email": "personal@example.com",
        }],
        column_headers=["企业邮箱", "个人邮箱"],
    )

    sheet = openpyxl.load_workbook(io.BytesIO(content)).active

    assert [cell.value for cell in sheet[1]] == ["企业邮箱", "个人邮箱"]
    assert [cell.value for cell in sheet[2]] == ["work@example.com", "personal@example.com"]


@pytest.mark.parametrize(("values", "expected"), [
    ([0.1, 0.2], 0.3),
    ([100.1, 200.2], 300.3),
    ([0.1] * 100, 10.0),
    ([0.3, -0.1, -0.2], 0.0),
    ([0.000001, 0.000002], 0.000003),
    ([9007199254740993, 2], 9007199254740995),
    (["100.10", "200.20"], 300.3),
    ([True, 0.2], 1.2),
])
def test_aggregate_records_preserves_decimal_sum(values, expected):
    rows, anomalies = engine.aggregate_records(
        [({"员工": "E001", "金额": value}, "来源") for value in values],
        ["员工"], ["金额"],
    )

    assert anomalies == []
    assert rows[0]["金额"] == expected
    if isinstance(expected, int):
        assert isinstance(rows[0]["金额"], int)
    assert json.loads(json.dumps(rows))[0]["金额"] == expected


@pytest.mark.parametrize("separate_sources", [False, True])
def test_merge_decimal_preview_matches_export(separate_sources):
    files = []
    mappings = []
    for index, values in enumerate(([100.1, 0.2], [200.2])):
        source = f"来源{index}" if separate_sources else "来源"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = source
        sheet.append(["员工", "金额", "备注"])
        for value in values:
            sheet.append(["E001", value, "-"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        files.append((f"{index}.xlsx", buffer.getvalue()))
        if separate_sources or not mappings:
            mappings.append({
                "name": source,
                "sheet_kw": source,
                "header": [1, 1],
                "match": ["员工", "金额", "备注"],
                "key_map": {"员工": "员工"},
                "column_map": {"金额": "金额"},
                "derived_fields": [],
            })

    result = engine.run_merge(
        files,
        {"merge_keys": ["员工"], "std_fields": ["金额"], "aggregate": "sum"},
        mappings,
    )
    preview = json.loads(json.dumps(result))["rows"][0]
    exported = openpyxl.load_workbook(io.BytesIO(
        engine.rows_to_xlsx(result["columns"], result["rows"])
    ), data_only=True)

    assert result["anomalies"] == []
    assert preview["金额"] == 300.5
    assert exported.active.cell(2, 2).value == preview["金额"]
